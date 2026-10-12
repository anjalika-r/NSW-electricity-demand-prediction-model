"""
NSW electricity demand — LONG-TERM forecast (any day, years into the future).

HOW THIS DIFFERS FROM forecast.py
---------------------------------
forecast.py predicts the NEXT HOUR using recent demand (lag_1, lag_24, lag_168).
Those lags are its engine. But for a day in 2030 there IS no "demand an hour ago",
so that model cannot see the future at all.

This model removes the lag features entirely and predicts each hour's demand from
ONLY things we can know for ANY date, past or future:
  - calendar   : hour, day-of-week, month, season  (known for any date, forever)
  - weather    : a TYPICAL year built from historical averages (a stand-in, since
                 real future weather is unknowable)
  - drivers    : NSW population and GSP for that year, from projections / growth
With the lags gone, the yearly drivers are what carry growth into future years.

WHAT IT GIVES YOU
-----------------
The EXPECTED demand shape and level for any day you choose, under your scenario
assumptions. It will NOT reproduce a specific day's weather spike years out
(nobody can) — it gives the typical profile scaled by that year's population/economy.

Data used (same folder as forecast.py, plus the two xlsx files):
  - PRICE_AND_DEMAND_*_NSW1.csv           AEMO 5-min demand (history, to train on)
  - Sydney.*.csv / Wollongong.*.csv       BoM daily weather (history)
  - 2024-nsw-population-projections-state.xlsx   NSW population, historic + to 2066
  - 5220001_Annual_Gross_State_Product_All_States.xlsx   NSW real GSP, to FY2025

Usage:
    pip install pandas numpy scikit-learn lightgbm openpyxl
    # train and check accuracy, then forecast a single future day:
    python longterm_forecast.py --data-dir ./data --forecast-date 2030-03-14
    # forecast a whole future year (hourly), save to CSV:
    python longterm_forecast.py --data-dir ./data --forecast-year 2030 --out year2030.csv

Assumptions you can change:
    --gsp-growth 0.02      annual real GSP growth past the last actual year (2% default)
"""

import argparse
import glob
import os
import re
import numpy as np
import pandas as pd


# ======================================================================
# 1. LOAD HISTORICAL DEMAND  (same logic as forecast.py)
# ======================================================================
def load_demand(data_dir):
    files = sorted(glob.glob(os.path.join(data_dir, "PRICE_AND_DEMAND_*_NSW1.csv")))
    if not files:
        raise FileNotFoundError("No PRICE_AND_DEMAND_*_NSW1.csv files found.")
    frames = [pd.read_csv(f, usecols=["SETTLEMENTDATE", "TOTALDEMAND"]) for f in files]
    demand = pd.concat(frames, ignore_index=True)
    demand["SETTLEMENTDATE"] = pd.to_datetime(
        demand["SETTLEMENTDATE"], format="%Y/%m/%d %H:%M:%S"
    )
    demand = demand.sort_values("SETTLEMENTDATE").drop_duplicates("SETTLEMENTDATE")

    # 5-min -> hourly. Timestamps mark the END of each interval, so shift back
    # 5 min then floor, so 00:05..01:00 count as hour 00.
    idx = (demand["SETTLEMENTDATE"] - pd.Timedelta(minutes=5)).dt.floor("h")
    hourly = (
        demand.assign(hour=idx)
        .groupby("hour")
        .agg(demand_mw=("TOTALDEMAND", "mean"))
        .reset_index()
        .rename(columns={"hour": "ts"})
    )
    return hourly


# ======================================================================
# 2. LOAD HISTORICAL WEATHER  (same parsing as forecast.py)
# ======================================================================
def _find_header_line(path):
    """Find the real BoM table header: a line that mentions Date AND is a genuine
    multi-column row (several commas). This skips preamble lines that merely
    contain the word 'Date' in prose."""
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        lines = fh.readlines()
    # 1st choice: a Date line that has at least a few comma-separated fields
    for i, line in enumerate(lines):
        if re.search(r'\bDate\b', line, re.IGNORECASE) and line.count(",") >= 3:
            return i
    # 2nd choice: any line starting with an optional blank col then "Date"
    for i, line in enumerate(lines):
        if re.match(r'^,?\s*"?Date"?\s*,', line, re.IGNORECASE):
            return i
    raise ValueError(f"No Date header found in {path}")


def load_weather(data_dir, prefix, label):
    patterns = [os.path.join(data_dir, f"{prefix}.*.csv")]
    if prefix == "Sydney":
        patterns.append(os.path.join(data_dir, "IDCJDW2124.*.csv"))  # misnamed May 2026 file
    files = sorted(set(sum([glob.glob(p) for p in patterns], [])))
    if not files:
        return pd.DataFrame()

    frames = []
    for f in files:
        skip = _find_header_line(f)
        df = pd.read_csv(f, skiprows=skip, encoding="utf-8", encoding_errors="ignore")
        df.columns = [str(c).strip() for c in df.columns]
        df = df.loc[:, ~df.columns.str.match(r"^Unnamed")]

        def col(sub):
            hits = [c for c in df.columns if sub.lower() in c.lower()]
            return hits[0] if hits else None

        out = pd.DataFrame()
        out["date"] = pd.to_datetime(df[col("Date")], errors="coerce")
        for sub, name in {
            "Minimum temperature": "tmin",
            "Maximum temperature": "tmax",
            "Rainfall": "rain",
            "9am Temperature": "t9",
            "9am relative humidity": "rh9",
            "3pm Temperature": "t3",
            "3pm relative humidity": "rh3",
        }.items():
            c = col(sub)
            out[name] = pd.to_numeric(df[c], errors="coerce") if c else np.nan
        out = out.dropna(subset=["date"])
        frames.append(out)

    w = pd.concat(frames, ignore_index=True).drop_duplicates("date").sort_values("date")
    return w.add_suffix(f"_{label}").rename(columns={f"date_{label}": "date"})


# ======================================================================
# 3. LOAD ANNUAL DRIVERS  (population, GSP)  -> one value per year
# ======================================================================
def load_population(xlsx_path):
    """NSW total population by year (historic + projected to 2066)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = list(wb["Total population"].iter_rows(values_only=True))
    years, vals = rows[6], rows[7]            # row 6 = year headers, row 7 = NSW totals
    pop = {y: v for y, v in zip(years, vals) if isinstance(y, int)}
    return pd.Series(pop, name="population")


def load_gsp(xlsx_path, growth, horizon_year):
    """NSW real GSP by year. ABS stops at FY2025; extend to horizon_year at `growth`.

    Each ABS annual obs is a financial year ending 30 June -> we key it by that
    calendar year (FY2024-25 -> 2025). Good enough for an annual driver.
    """
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = list(wb["Data1"].iter_rows(values_only=True))
    gsp = {}
    for r in rows[10:]:                        # data starts at row 10
        if r[0] is not None and r[1] is not None:
            gsp[r[0].year] = float(r[1])       # column 1 = New South Wales

    last_year = max(gsp)
    last_val = gsp[last_year]
    # extend forward at the assumed growth rate (DOCUMENT THIS ASSUMPTION)
    for y in range(last_year + 1, horizon_year + 1):
        last_val *= (1 + growth)
        gsp[y] = last_val
    s = pd.Series(gsp, name="gsp")
    s.attrs["last_actual_year"] = last_year    # remember where real data ends
    return s


# ======================================================================
# 4. TYPICAL-YEAR WEATHER  (the stand-in for unknowable future weather)
# ======================================================================
def build_typical_weather(syd, wol):
    """Average each weather field by (month, day) across all historical years.
    This is the 'typical year' a future date borrows its weather from.
    Returns a frame keyed by (month, day)."""
    frames = []
    for w, lab in ((syd, "syd"), (wol, "wol")):
        if w.empty:
            continue
        w = w.copy()
        w["month"] = w["date"].dt.month
        w["day"] = w["date"].dt.day
        cols = [c for c in w.columns if c not in ("date", "month", "day")]
        typ = w.groupby(["month", "day"])[cols].mean()
        frames.append(typ)
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, axis=1).reset_index()


# ======================================================================
# 5. FEATURES  — ONLY future-knowable inputs. NO demand lags.
# ======================================================================
def calendar_features(df):
    ts = df["ts"]
    df["hour"] = ts.dt.hour
    df["dow"] = ts.dt.dayofweek
    df["month"] = ts.dt.month
    df["year"] = ts.dt.year
    df["is_weekend"] = (df["dow"] >= 5).astype(int)
    df["hour_sin"] = np.sin(2 * np.pi * df["hour"] / 24)
    df["hour_cos"] = np.cos(2 * np.pi * df["hour"] / 24)
    df["dow_sin"] = np.sin(2 * np.pi * df["dow"] / 7)
    df["dow_cos"] = np.cos(2 * np.pi * df["dow"] / 7)
    df["doy_sin"] = np.sin(2 * np.pi * ts.dt.dayofyear / 365)
    df["doy_cos"] = np.cos(2 * np.pi * ts.dt.dayofyear / 365)
    return df


def attach_drivers(df, pop, gsp):
    """Join the year's population and GSP onto each row by calendar year.
    No leakage shift needed: these are known years ahead (unlike daily weather)."""
    df["population"] = df["year"].map(pop)
    df["gsp"] = df["year"].map(gsp)
    return df


def build_training_features(hourly, syd, wol, pop, gsp):
    """Historical rows: real weather, real drivers, real demand (the target)."""
    df = hourly.copy()
    df = calendar_features(df)
    df["date"] = df["ts"].dt.normalize()

    # real historical daily weather, attached to each hour of that day
    for w in (syd, wol):
        if not w.empty:
            df = df.merge(w, on="date", how="left")

    df = attach_drivers(df, pop, gsp)
    return df


FEATURE_COLS = [
    "hour", "dow", "month", "is_weekend",
    "hour_sin", "hour_cos", "dow_sin", "dow_cos", "doy_sin", "doy_cos",
    "population", "gsp",
    # weather (typical-year or historical); missing cols are dropped automatically
    "tmax_syd", "tmin_syd", "t9_syd", "t3_syd", "rh9_syd", "rh3_syd", "rain_syd",
    "tmax_wol", "tmin_wol", "t9_wol", "t3_wol", "rh9_wol", "rh3_wol", "rain_wol",
]


# ======================================================================
# 6. MODEL
# ======================================================================
def make_model():
    try:
        from lightgbm import LGBMRegressor
        return LGBMRegressor(
            n_estimators=800, learning_rate=0.03, num_leaves=64,
            subsample=0.8, colsample_bytree=0.8, random_state=42,
        )
    except ImportError:
        from sklearn.ensemble import HistGradientBoostingRegressor
        return HistGradientBoostingRegressor(
            max_iter=800, learning_rate=0.03, random_state=42
        )


# Driver columns are handled by the LINEAR level model (they extrapolate there),
# NOT by the tree — a tree clamps them at the training range and flatlines.
DRIVER_COLS = ["population", "gsp"]


def shape_features(df):
    """Feature cols for the TREE: everything except the annual drivers."""
    return [c for c in FEATURE_COLS if c in df.columns and c not in DRIVER_COLS]


def present_features(df):
    return [c for c in FEATURE_COLS if c in df.columns]


# ----------------------------------------------------------------------
# Two-layer model (Option B): linear LEVEL on drivers + tree SHAPE.
#
#   level(year)  = a linear fit of ANNUAL-MEAN demand on population & GSP.
#                  Linear => it extrapolates: more people -> proportionally
#                  more demand, even past the training range.
#   shape(hour)  = a tree fit on (actual demand - that year's mean), using only
#                  calendar + weather. Captures the daily/seasonal pattern.
#   prediction   = level(year) + shape(hour)
# ----------------------------------------------------------------------
class TwoLayerModel:
    def __init__(self):
        self.level = None          # LinearRegression on drivers
        self.shape = None          # tree on calendar+weather
        self.shape_cols = None

    def fit(self, feats, pop, gsp):
        from sklearn.linear_model import LinearRegression

        feats = feats.dropna(subset=["demand_mw"]).copy()

        # --- LEVEL: one row per year (year-mean demand vs that year's drivers)
        yearly = (feats.groupby("year")
                  .agg(demand_mw=("demand_mw", "mean"))
                  .reset_index())
        yearly["population"] = yearly["year"].map(pop)
        yearly["gsp"] = yearly["year"].map(gsp)
        self.level = LinearRegression().fit(
            yearly[DRIVER_COLS].values, yearly["demand_mw"].values
        )
        self._yearly = yearly      # kept for reporting

        # --- SHAPE: tree on the residual after removing each year's mean
        year_mean = feats["year"].map(
            feats.groupby("year")["demand_mw"].mean()
        )
        resid = feats["demand_mw"].values - year_mean.values
        self.shape_cols = shape_features(feats)
        self.shape = make_model().fit(feats[self.shape_cols], resid)
        return self

    def predict(self, df):
        lvl = self.level.predict(df[DRIVER_COLS].values)
        shp = self.shape.predict(df[self.shape_cols])
        return lvl + shp

    def level_for_year(self, pop_val, gsp_val):
        import numpy as np
        return float(self.level.predict(np.array([[pop_val, gsp_val]]))[0])


def train_and_eval(feats, pop, gsp):
    from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

    feats = feats.dropna(subset=["demand_mw"]).sort_values("ts").reset_index(drop=True)

    cut = int(len(feats) * 0.8)                 # chronological split
    tr, te = feats.iloc[:cut], feats.iloc[cut:]

    model = TwoLayerModel().fit(tr, pop, gsp)
    pred = model.predict(te)
    yte = te["demand_mw"].values

    print("\n=== Backtest (chronological last 20%) ===")
    print(f"rows : {len(yte)}")
    print(f"MAE  : {mean_absolute_error(yte, pred):8.1f} MW")
    print(f"RMSE : {np.sqrt(mean_squared_error(yte, pred)):8.1f} MW")
    print(f"MAPE : {np.mean(np.abs((yte - pred) / yte)) * 100:8.2f} %")
    print(f"R^2  : {r2_score(yte, pred):8.4f}")

    # Report what the LINEAR level model learned — this is the slope being
    # extrapolated to 2060, so look at it critically.
    coef_pop, coef_gsp = model.level.coef_
    print("\n--- Linear level model (what drives the long-run trend) ---")
    print(f"  demand_mean = {model.level.intercept_:.0f}"
          f"  + {coef_pop:.6f} x population"
          f"  + {coef_gsp:.6f} x GSP")
    print(f"  => +1,000,000 people  ~  {coef_pop*1e6:+.0f} MW mean demand")
    print(f"  => +10,000 $m GSP     ~  {coef_gsp*1e4:+.0f} MW mean demand")
    n_years = model._yearly["year"].nunique()
    print(f"  fitted on {n_years} yearly data point(s)"
          + ("  <-- VERY FEW: slope is fragile, state this in your report"
             if n_years < 5 else ""))
    print("(No demand lags, so hour-to-hour accuracy is weaker than forecast.py —")
    print(" this model targets the typical PROFILE scaled by the driver trend.)")

    # refit on ALL history for actual forecasting
    model = TwoLayerModel().fit(feats, pop, gsp)
    return model, model.shape_cols


# ======================================================================
# 7. BUILD A FUTURE DAY / YEAR AND PREDICT
# ======================================================================
def future_frame(dates_index, typ_weather, pop, gsp):
    """Make the feature rows for a set of future hourly timestamps."""
    df = pd.DataFrame({"ts": dates_index})
    df = calendar_features(df)
    df["day"] = df["ts"].dt.day

    if not typ_weather.empty:
        df = df.merge(typ_weather, on=["month", "day"], how="left")

    df = attach_drivers(df, pop, gsp)
    return df


def _check_drivers(df, pop, gsp):
    yr = int(df["year"].iloc[0])
    if yr not in pop.index:
        raise ValueError(f"No population value for {yr} (have up to {pop.index.max()}).")
    if yr not in gsp.index:
        raise ValueError(f"No GSP value for {yr} (extend --gsp-growth horizon).")


def plot_day(df, date_str, path="forecast_day.png"):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(11, 5))
    ax.plot(df["hour"], df["pred_mw"], marker="o", color="steelblue", lw=2)
    peak = df.loc[df["pred_mw"].idxmax()]
    ax.scatter([peak["hour"]], [peak["pred_mw"]], color="crimson", zorder=5)
    ax.annotate(f"peak {peak['pred_mw']:.0f} MW",
                (peak["hour"], peak["pred_mw"]),
                textcoords="offset points", xytext=(0, 10), color="crimson")
    ax.set_title(f"NSW demand forecast — {date_str} (typical-weather scenario)")
    ax.set_xlabel("hour of day")
    ax.set_ylabel("demand (MW)")
    ax.set_xticks(range(0, 24, 2))
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    print(f"saved day plot -> {path}")


def plot_year(df, year, path="forecast_year.png"):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    daily = df.set_index("ts")["pred_mw"].resample("D").agg(["mean", "max", "min"])

    fig, ax = plt.subplots(2, 1, figsize=(13, 8), sharex=True)
    # top: daily mean with min-max band — shows the seasonal shape
    ax[0].fill_between(daily.index, daily["min"], daily["max"],
                       color="steelblue", alpha=0.2, label="daily min–max")
    ax[0].plot(daily.index, daily["mean"], color="steelblue", lw=1.2, label="daily mean")
    ax[0].set_title(f"NSW demand forecast — {year} (typical-weather scenario)")
    ax[0].set_ylabel("demand (MW)")
    ax[0].legend(loc="upper right")
    ax[0].grid(alpha=0.3)

    # bottom: one example week (mid-July) so the daily pattern is visible
    wk = df[(df["ts"] >= f"{year}-07-14") & (df["ts"] < f"{year}-07-21")]
    ax[1].plot(wk["ts"], wk["pred_mw"], color="darkorange", lw=1.5)
    ax[1].set_title("Example week (14–20 July) — hourly detail")
    ax[1].set_ylabel("demand (MW)")
    ax[1].set_xlabel("date")
    ax[1].grid(alpha=0.3)

    fig.tight_layout()
    fig.savefig(path, dpi=120)
    print(f"saved year plot -> {path}")


def forecast_day(model, cols, date_str, typ_weather, pop, gsp, plot=True):
    day = pd.Timestamp(date_str)
    idx = pd.date_range(day, day + pd.Timedelta(hours=23), freq="h")
    df = future_frame(idx, typ_weather, pop, gsp)
    _check_drivers(df, pop, gsp)
    df["pred_mw"] = model.predict(df)

    print(f"\n=== Forecast for {date_str} (typical-weather scenario) ===")
    print(f"population used : {int(df['population'].iloc[0]):,}")
    print(f"GSP used        : {df['gsp'].iloc[0]:,.0f} $m")
    print(f"daily peak      : {df['pred_mw'].max():8.1f} MW  at {df.loc[df['pred_mw'].idxmax(),'hour']:02d}:00")
    print(f"daily min       : {df['pred_mw'].min():8.1f} MW")
    print(f"daily mean      : {df['pred_mw'].mean():8.1f} MW")
    print(f"daily energy    : {df['pred_mw'].sum()/1000:8.1f} GWh")
    if plot:
        plot_day(df, date_str)
    return df[["ts", "pred_mw"]]


def forecast_year(model, cols, year, typ_weather, pop, gsp, out=None, plot=True):
    idx = pd.date_range(f"{year}-01-01", f"{year}-12-31 23:00", freq="h")
    df = future_frame(idx, typ_weather, pop, gsp)
    _check_drivers(df, pop, gsp)
    df["pred_mw"] = model.predict(df)

    print(f"\n=== Forecast for all of {year} (typical-weather scenario) ===")
    print(f"population used : {int(df['population'].iloc[0]):,}")
    print(f"GSP used        : {df['gsp'].iloc[0]:,.0f} $m")
    print(f"annual energy   : {df['pred_mw'].sum()/1e3:,.0f} GWh")
    print(f"annual peak     : {df['pred_mw'].max():8.1f} MW")
    print(f"annual min      : {df['pred_mw'].min():8.1f} MW")
    out_df = df[["ts", "pred_mw"]]
    if out:
        out_df.to_csv(out, index=False)
        print(f"saved hourly forecast -> {out}")
    if plot:
        plot_year(df, year)
    return out_df


# Generation-design capacities (MW). Edit the numbers to match your designs.
# (name, capacity_MW, colour)
CAPACITY_THRESHOLDS = [
    ("Solar", 15000, "orange"),
    ("Gas",   13000, "firebrick"),
    ("Wind",  12000, "seagreen"),
]


def forecast_trend(model, cols, start, end, typ_weather, pop, gsp,
                   out=None, path="demand_trend.png"):
    """Average demand per year across a range of years — the long-run trend.
    For each year, build a full typical year, predict every hour, take the mean."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    years, avg_mw, peak_mw, energy_gwh = [], [], [], []
    for yr in range(start, end + 1):
        if yr not in pop.index or yr not in gsp.index:
            print(f"  skipping {yr} (no driver data)")
            continue
        idx = pd.date_range(f"{yr}-01-01", f"{yr}-12-31 23:00", freq="h")
        df = future_frame(idx, typ_weather, pop, gsp)
        p = model.predict(df)
        years.append(yr)
        avg_mw.append(p.mean())
        peak_mw.append(p.max())
        energy_gwh.append(p.sum() / 1e3)

    trend = pd.DataFrame(
        {"year": years, "avg_mw": avg_mw, "peak_mw": peak_mw, "energy_gwh": energy_gwh}
    )

    print(f"\n=== Average demand per year, {start}-{end} ===")
    print(trend.to_string(index=False,
          formatters={"avg_mw": "{:.0f}".format, "peak_mw": "{:.0f}".format,
                      "energy_gwh": "{:,.0f}".format}))

    pct = (trend["avg_mw"].iloc[-1] / trend["avg_mw"].iloc[0] - 1) * 100
    print(f"\naverage demand grows {pct:+.1f}% from {years[0]} to {years[-1]}")

    fig, ax = plt.subplots(figsize=(12, 6))
    ax.plot(trend["year"], trend["avg_mw"], marker="o", lw=2, color="steelblue",
            label="average demand", zorder=3)
    ax.fill_between(trend["year"], trend["avg_mw"], alpha=0.08, color="steelblue")
    ax.set_title(f"NSW average electricity demand per year, {start}–{end}\n"
                 f"(typical-weather scenario, GSP growth as assumed)")
    ax.set_xlabel("year")
    ax.set_ylabel("average demand (MW)")
    ax.grid(alpha=0.3)
    # annotate first and last points
    for i in (0, -1):
        ax.annotate(f"{trend['avg_mw'].iloc[i]:.0f} MW",
                    (trend["year"].iloc[i], trend["avg_mw"].iloc[i]),
                    textcoords="offset points", xytext=(0, 10),
                    ha="center", color="black")

    # --- capacity thresholds: when does demand reach each design's capacity? ---
    yvals = trend["avg_mw"].values
    xvals = trend["year"].values
    y0 = ax.get_ylim()[0]
    x0 = xvals[0]
    for name, cap, colour in CAPACITY_THRESHOLDS:
        # find the first year demand crosses this capacity (linear interpolation)
        cross_year = None
        for i in range(1, len(yvals)):
            if yvals[i - 1] < cap <= yvals[i]:
                frac = (cap - yvals[i - 1]) / (yvals[i] - yvals[i - 1])
                cross_year = xvals[i - 1] + frac * (xvals[i] - xvals[i - 1])
                break

        # legend handle (drawn off-canvas so the label still shows)
        ax.plot([], [], color=colour, ls=":", lw=1.8,
                label=f"{name} capacity ({cap:,} MW)")

        if cross_year is None:
            where = "below range" if cap > yvals[-1] else "already exceeded"
            print(f"  {name}: {cap:,} MW not crossed in {start}-{end} ({where})")
            continue
        # horizontal segment: from left edge ONLY up to where it meets the curve
        ax.hlines(cap, x0, cross_year, color=colour, ls=":", lw=1.8, zorder=2)
        # vertical segment: from the crossing point DOWN to the x-axis
        ax.vlines(cross_year, y0, cap, color=colour, ls=":", lw=1.8, zorder=2)
        # marker at the meeting point on the demand curve
        ax.plot([cross_year], [cap], marker="o", color=colour, ms=8,
                zorder=4, markeredgecolor="white")
        ax.annotate(f"{cross_year:.0f}", (cross_year, y0),
                    textcoords="offset points", xytext=(0, 4),
                    ha="center", va="bottom", color=colour, fontweight="bold")
        print(f"  {name:6s}: {cap:,} MW reached in {cross_year:.1f}")

    ax.set_xlim(left=x0)
    ax.set_ylim(bottom=y0)
    ax.legend(loc="upper left", framealpha=0.9)
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    print(f"saved trend plot -> {path}")

    if out:
        trend.to_csv(out, index=False)
        print(f"saved trend data -> {out}")
    return trend


# ======================================================================
# 8. MAIN
# ======================================================================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=".")
    ap.add_argument("--pop-file", default="2024-nsw-population-projections-state.xlsx")
    ap.add_argument("--gsp-file", default="5220001_Annual_Gross_State_Product_All_States.xlsx")
    ap.add_argument("--gsp-growth", type=float, default=0.02,
                    help="assumed annual real GSP growth past last actual year")
    ap.add_argument("--forecast-date", help="YYYY-MM-DD, forecast a single day")
    ap.add_argument("--forecast-year", type=int, help="YYYY, forecast the whole year")
    ap.add_argument("--trend", help="START-END, e.g. 2026-2060: average demand per year")
    ap.add_argument("--out", help="CSV path for --forecast-year / --trend output")
    args = ap.parse_args()

    # how far ahead the drivers must reach
    trend_end = int(args.trend.split("-")[1]) if args.trend else 0
    horizon = max(
        trend_end,
        args.forecast_year or 0,
        pd.Timestamp(args.forecast_date).year if args.forecast_date else 0,
    ) or 2035

    print("Loading demand history...")
    hourly = load_demand(args.data_dir)
    print(f"  {len(hourly)} hourly rows, {hourly['ts'].min()} -> {hourly['ts'].max()}")

    print("Loading weather history...")
    syd = load_weather(args.data_dir, "Sydney", "syd")
    wol = load_weather(args.data_dir, "Wollongong", "wol")

    print("Loading annual drivers...")
    pop = load_population(os.path.join(args.data_dir, args.pop_file))
    gsp = load_gsp(os.path.join(args.data_dir, args.gsp_file), args.gsp_growth, horizon)
    print(f"  population {pop.index.min()}-{pop.index.max()}; "
          f"GSP actual to {gsp.attrs['last_actual_year']}, "
          f"extended to {gsp.index.max()} at {args.gsp_growth:.0%}/yr")

    print("Building typical-year weather...")
    typ = build_typical_weather(syd, wol)

    print("Training...")
    feats = build_training_features(hourly, syd, wol, pop, gsp)
    model, cols = train_and_eval(feats, pop, gsp)
    print(f"\nfeatures used: {cols}")

    if args.forecast_date:
        forecast_day(model, cols, args.forecast_date, typ, pop, gsp)
    if args.forecast_year:
        forecast_year(model, cols, args.forecast_year, typ, pop, gsp, args.out)
    if args.trend:
        start, end = (int(x) for x in args.trend.split("-"))
        forecast_trend(model, cols, start, end, typ, pop, gsp, args.out)
    if not (args.forecast_date or args.forecast_year or args.trend):
        print("\nTip: --forecast-date 2030-03-14 | --forecast-year 2030 | --trend 2026-2060")


if __name__ == "__main__":
    main()